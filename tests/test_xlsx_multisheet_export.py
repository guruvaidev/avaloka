"""
tests/test_xlsx_multisheet_export.py

Multi-sheet Excel uploads become one dataset per sheet.

Follow-up to test_excel_sheet_selection.py: picking a better single sheet is
not enough — the client expectation is that Avaloka knows about EVERY sheet
in an uploaded workbook. /api/upload now expands a workbook into sibling
datasets (one per non-empty sheet), each materialized as a CSV so the whole
downstream pipeline (sampling, codegen, execution) runs on the CSV-native
path.

Unit tests here need only pandas + openpyxl. The /api/upload tests reuse the
fakes from test_server_unit and are skipped automatically when the
server's dependency stack (fastapi, daft, langchain, ...) is not importable.
"""

import csv
import io
import re

import pandas as pd
import pytest
from openpyxl import Workbook

from file_handler.excel_connector import export_workbook_sheets_to_csv


# ---------------------------------------------------------------------------
# Fixtures / helpers
# ---------------------------------------------------------------------------

def _df(rows: int, prefix: str = "") -> pd.DataFrame:
    return pd.DataFrame(
        {
            f"{prefix}id": range(1, rows + 1),
            f"{prefix}amount": [round(10.5 * i, 2) for i in range(1, rows + 1)],
            f"{prefix}category": ["alpha" if i % 2 else "beta" for i in range(rows)],
        }
    )


def _write_workbook(path, sheets: dict) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)


def _workbook_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    _write_workbook(buf, sheets)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Unit: export_workbook_sheets_to_csv
# ---------------------------------------------------------------------------

def test_export_orders_largest_sheet_first_and_skips_empty(tmp_path):
    wb = tmp_path / "wb.xlsx"
    _write_workbook(
        wb,
        {
            "Notes": pd.DataFrame({"note": ["cover", "decoy", "text"]}),
            "Data": _df(20),
            "Blank": pd.DataFrame(),
        },
    )

    exports, notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    assert [e.sheet_name for e in exports] == ["Data", "Notes"]
    assert exports[0].n_rows == 20 and exports[1].n_rows == 3

    # Sheet CSV round-trips with columns and dtypes usable downstream.
    out = pd.read_csv(exports[0].csv_path)
    assert list(out.columns) == ["id", "amount", "category"]
    assert len(out) == 20
    assert pd.api.types.is_numeric_dtype(out["amount"])

    assert any("Blank" in n and "empty" in n.lower() for n in notes)


def test_export_max_sheets_keeps_largest_and_notes_dropped(tmp_path):
    wb = tmp_path / "wb.xlsx"
    _write_workbook(
        wb,
        {"S1": _df(2), "S2": _df(8), "S3": _df(5), "S4": _df(1)},
    )

    exports, notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out", max_sheets=2)

    assert [e.sheet_name for e in exports] == ["S2", "S3"]
    dropped_note = next(n for n in notes if "Not imported" in n)
    assert "S1" in dropped_note and "S4" in dropped_note
    # Dropped sheets leave no CSV files behind.
    remaining = {e.csv_path for e in exports}
    on_disk = {str(p) for p in (tmp_path / "out").glob("*.csv")}
    assert on_disk == remaining


def test_export_single_sheet_workbook(tmp_path):
    wb = tmp_path / "wb.xlsx"
    _write_workbook(wb, {"Only": _df(5)})

    exports, notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    assert len(exports) == 1
    assert exports[0].sheet_name == "Only"
    assert exports[0].n_rows == 5
    assert notes == []


def test_export_all_sheets_empty_returns_no_exports(tmp_path):
    wb = tmp_path / "wb.xlsx"
    _write_workbook(wb, {"A": pd.DataFrame(), "B": pd.DataFrame()})

    exports, notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    assert exports == []
    assert len(notes) == 2


# ---------------------------------------------------------------------------
# A sheet is a canvas: the table does not start at A1
# ---------------------------------------------------------------------------
# file_handler/table_shape.py finds where the table starts, and ExcelConnector
# has used it since it was written. This export did not, and it is the function
# /api/upload calls for every workbook -- so the capability existed, was tested
# and was documented, and a workbook uploaded through the UI never reached it.

CANVAS_HEADER = ["Run", "Start (s)", "End (s)", "Planned step", "Result", "Score (1-5)"]
CANVAS_NUMERIC = ["Run", "Start (s)", "End (s)", "Score (1-5)"]
CANVAS_ROWS = 12
_UNNAMED = re.compile(r"^Unnamed:\s*\d+$")


def _canvas_rows() -> list:
    return [
        [i, round(1.5 * i, 2), round(1.5 * i + 1.25, 2), f"step {i}", "done" if i % 2 else "partial", 1 + i % 5]
        for i in range(1, CANVAS_ROWS + 1)
    ]


def write_canvas_workbook(target, rows=None, blank_row_under_title: bool = True) -> None:
    """A table laid out the way people lay them out: column A left empty, a
    one-cell title on row 2, optionally a blank row, then the header, with the
    panes frozen under the header and beside the indent."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Annotations"
    sheet.cell(row=2, column=2, value="Synthetic annotation export")
    header_at = 4 if blank_row_under_title else 3
    for offset, name in enumerate(CANVAS_HEADER):
        sheet.cell(row=header_at, column=2 + offset, value=name)
    for r, record in enumerate(_canvas_rows() if rows is None else rows, start=header_at + 1):
        for offset, value in enumerate(record):
            sheet.cell(row=r, column=2 + offset, value=value)
    sheet.freeze_panes = sheet.cell(row=header_at + 1, column=2)
    book.save(target)


@pytest.mark.parametrize("blank_row_under_title", [True, False])
def test_export_finds_the_header_under_a_title_row_and_an_indent(tmp_path, blank_row_under_title):
    wb = tmp_path / "canvas.xlsx"
    write_canvas_workbook(wb, blank_row_under_title=blank_row_under_title)

    exports, notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    assert len(exports) == 1 and notes == []
    out = pd.read_csv(exports[0].csv_path)
    assert list(out.columns) == CANVAS_HEADER
    assert not [c for c in out.columns if _UNNAMED.match(str(c))]
    # The header is not one of the rows, so the columns under it can be typed.
    assert exports[0].n_rows == CANVAS_ROWS and len(out) == CANVAS_ROWS
    for name in CANVAS_NUMERIC:
        assert pd.api.types.is_numeric_dtype(out[name]), (name, out[name].dtype)
    assert out["Start (s)"].tolist() == [round(1.5 * i, 2) for i in range(1, CANVAS_ROWS + 1)]
    assert out["Planned step"].tolist() == [f"step {i}" for i in range(1, CANVAS_ROWS + 1)]


def test_null_word_labels_survive_the_export_of_an_indented_table(tmp_path):
    """The words are read from the whole sheet and the table is cut out of it,
    so they have to be cut out with it (_aligned_words) or they land a row and
    a column away from the cells they came from."""
    rows = _canvas_rows()
    rows[1][4] = "None"         # a label in a text column: kept
    rows[2][4] = "N/A"
    rows[3][5] = "N/A"          # a gap in a numeric column: still a gap
    wb = tmp_path / "canvas.xlsx"
    write_canvas_workbook(wb, rows=rows)

    (export,), _notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    with open(export.csv_path, newline="") as handle:
        read = list(csv.DictReader(handle))
    assert [r["Result"] for r in read[:4]] == ["done", "None", "N/A", "partial"]
    assert read[3]["Score (1-5)"] == ""
    assert pd.api.types.is_numeric_dtype(pd.read_csv(export.csv_path)["Score (1-5)"])


@pytest.mark.parametrize(
    "header",
    [
        ["Region", 2021, 2022, 2023],                                    # years
        ["team", "score", "score", "lead"],                              # a repeated name
        ["id", "How satisfied were you with the service you received?", "x", "y"],
    ],
    ids=["numeric-names", "repeated-name", "a-sentence"],
)
def test_a_sheet_that_starts_at_a1_is_exported_as_it_always_was(tmp_path, header):
    """Only a sheet showing the symptom is re-read. table_shape declines these
    three headers (too many numbers, a repeated label, prose) and would hand
    back column_1.. with the header as the first record; pandas reads them
    correctly today, and a fix for the canvas case must not cost that."""
    book = Workbook()
    sheet = book.active
    for record in [header, ["North", 1, 2, 3], ["South", 4, 5, 6]]:
        sheet.append(record)
    wb = tmp_path / "plain.xlsx"
    book.save(wb)

    (export,), _notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    out = pd.read_csv(export.csv_path)
    assert export.n_rows == 2 and len(out) == 2
    assert out.iloc[:, 0].tolist() == ["North", "South"]
    assert not [c for c in out.columns if re.match(r"^column_\d+$", str(c))]
    for position in (1, 2, 3):
        assert pd.api.types.is_numeric_dtype(out.iloc[:, position])


def test_a_sheet_with_two_tables_is_still_one_csv(tmp_path):
    """Decided, not overlooked: load_all_tables splits a sheet into its tables,
    this export does not. One sheet stays one dataset."""
    book = Workbook()
    sheet = book.active
    for r, record in enumerate([["id", "amount"], [1, 10.5], [2, 21.0], [3, 31.5]], start=1):
        for c, value in enumerate(record, start=1):
            sheet.cell(row=r, column=c, value=value)
    for r, record in enumerate([["team", "lead"], ["A", "Ann"], ["B", "Bo"], ["C", "Cy"]], start=1):
        for c, value in enumerate(record, start=4):      # column C left empty between them
            sheet.cell(row=r, column=c, value=value)
    wb = tmp_path / "two.xlsx"
    book.save(wb)

    exports, _notes = export_workbook_sheets_to_csv(str(wb), tmp_path / "out")

    assert len(exports) == 1 and exports[0].n_rows == 3
    assert {"id", "amount", "team", "lead"} <= set(pd.read_csv(exports[0].csv_path).columns)
