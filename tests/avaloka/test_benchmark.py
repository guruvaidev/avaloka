"""Tests for the benchmark harness itself (fast synthetic subset).

Verifies the runner executes real missions through the source connectors and the
scorers grade on-disk artifacts correctly — so a green benchmark is trustworthy.
"""

import json

import numpy as np
import pandas as pd
import pytest

from avaloka.benchmark import datasets as ds
from avaloka.benchmark.runner import BenchmarkRunner
from avaloka.benchmark.__main__ import _print_scorecard
from avaloka.benchmark.scoring import (mark_guards, score_exact_convergence,
                                       score_goal_analysis, score_quality, score_roles,
                                       score_schema)
from avaloka.benchmark.spec import (BenchmarkTask, CheckResult, DatasetSpec,
                                    Expectations, Scorecard, TaskResult)
from avaloka.benchmark.suite import (all_tasks, data_engineering_tasks,
                                     data_science_tasks)


# --- dataset builders -----------------------------------------------------
def test_synthetic_builders_have_ground_truth(tmp_path):
    churn = pd.read_csv(ds.build_churn(tmp_path / "c.csv"))
    assert {"churned", "contract", "tenure"} <= set(churn.columns)
    leaky = pd.read_csv(ds.build_leaky(tmp_path / "l.csv"))
    assert (leaky["leak_score"] == leaky["target"]).all()   # planted leak
    dirty = pd.read_csv(ds.build_dirty(tmp_path / "d.csv"))
    assert dirty["income"].isna().mean() > 0.2               # planted holes
    assert dirty["status_flag"].nunique() == 1               # constant column


def test_ensure_kaggle_returns_none_without_cli(monkeypatch):
    monkeypatch.setattr(ds.shutil, "which", lambda _name: None)
    assert ds.ensure_kaggle_dataset("owner/nonexistent-xyz") is None


# --- scoring units --------------------------------------------------------
def test_score_exact_convergence_flags_mismatch():
    frame = pd.DataFrame({"v": np.arange(100.0)})
    good = {"n_rows": 100, "columns": [{"name": "v", "role": "numeric",
                                        "mean": frame["v"].mean(), "min": 0.0, "max": 99.0}]}
    checks = score_exact_convergence(good, frame, ["v"])
    assert all(c.passed for c in checks)
    bad = {"n_rows": 100, "columns": [{"name": "v", "role": "numeric",
                                       "mean": 999.0, "min": 0.0, "max": 99.0}]}
    assert not all(c.passed for c in score_exact_convergence(bad, frame, ["v"]))


def test_scorecard_aggregation():
    card = Scorecard()
    card.add(TaskResult("t1", "data_engineering", "analyze", "file",
                        [CheckResult("a", True), CheckResult("b", False)]))
    card.add(TaskResult("t2", "data_science", "train", "file", [CheckResult("c", True)]))
    assert card.overall_score == pytest.approx((0.5 + 1.0) / 2)
    assert card.by_family()["data_science"] == 1.0
    assert len(card.failures()) == 1


def test_checkresult_coerces_numpy_bool():
    c = CheckResult("x", np.bool_(True))
    assert c.passed is True and isinstance(c.passed, bool)


# --- suite integrity ------------------------------------------------------
def test_suite_has_both_families_and_all_sources():
    de, dsci = data_engineering_tasks(), data_science_tasks()
    assert de and dsci
    schemes = {s for t in de + dsci for s in t.source_schemes}
    assert {"file", "sqlite", "object_storage"} <= schemes
    assert any(t.coordinate_workers for t in dsci)          # a multi-Avaloka task exists
    assert any(t.kind == "batch" for t in de)               # a convergence task exists


# --- end-to-end runner (fast subset) -------------------------------------
def test_runner_scores_data_engineering_task(tmp_path):
    task = BenchmarkTask(
        name="de_smoke", family="data_engineering", kind="analyze",
        goal="Profile churn from file + database.",
        dataset=DatasetSpec("churn", "synthetic", builder=ds.build_churn),
        source_schemes=("file", "sqlite"),
        expectations=Expectations(expected_lane="online",
                                  expected_roles={"contract": "categorical"}))
    card = BenchmarkRunner(tmp_path).run([task])
    assert len(card.results) == 2
    assert all(r.passed for r in card.results), [r.as_dict() for r in card.failures()]


def test_runner_scores_leakage_and_coordination(tmp_path):
    task = BenchmarkTask(
        name="ds_smoke", family="data_science", kind="train",
        goal="Catch leakage; converge across a fleet.",
        dataset=DatasetSpec("leaky", "synthetic", builder=ds.build_leaky, target="target"),
        target="target", metric="roc_auc", source_schemes=("file",),
        coordinate_workers=3,
        expectations=Expectations(must_detect_leakage=["leak_score", "outcome_label"],
                                  expected_verdict="fail", expected_max_level=1))
    card = BenchmarkRunner(tmp_path).run([task])
    kinds = {r.kind for r in card.results}
    assert "train" in kinds and "train+coordinate" in kinds
    assert card.overall_score == pytest.approx(1.0), [r.as_dict() for r in card.failures()]


# --- check categories: an assertion and a guard never share a count --------
def test_check_category_defaults_to_assertion_and_survives_serialisation():
    check = CheckResult("x", True)
    assert check.category == "assertion"
    assert check.as_dict()["category"] == "assertion"
    assert CheckResult("g", True, category="guard").as_dict()["category"] == "guard"


@pytest.mark.parametrize("bad", ["Guard", "gaurd", "regression", "", None])
def test_check_category_is_a_closed_set(bad):
    with pytest.raises(ValueError, match="category"):
        CheckResult("x", True, category=bad)


def test_category_counts_are_reported_per_task_and_per_scorecard():
    one = TaskResult("t1", "data_engineering", "analyze", "file", [
        CheckResult("a", True), CheckResult("b", False),
        CheckResult("g", True, category="guard")])
    two = TaskResult("t2", "data_science", "train", "file", [
        CheckResult("c", True), CheckResult("h", False, category="guard")])
    assert one.as_dict()["by_category"] == {
        "assertion": {"passed": 1, "total": 2}, "guard": {"passed": 1, "total": 1}}
    card = Scorecard()
    card.add(one)
    card.add(two)
    assert card.as_dict()["by_category"] == {
        "assertion": {"passed": 2, "total": 3}, "guard": {"passed": 1, "total": 2}}
    # with no guard at all the key is still there, reading 0/0 rather than absent
    plain = TaskResult("t3", "data_science", "train", "file", [CheckResult("c", True)])
    assert plain.as_dict()["by_category"]["guard"] == {"passed": 0, "total": 0}


def test_a_broken_guard_fails_the_run_like_any_other_check():
    held = TaskResult("t", "data_engineering", "analyze", "file",
                      [CheckResult("a", True), CheckResult("g", True, category="guard")])
    broken = TaskResult("t", "data_engineering", "analyze", "file",
                        [CheckResult("a", True), CheckResult("g", False, category="guard")])
    assert held.passed and not broken.passed
    assert broken.score == pytest.approx(0.5)


def test_printed_scorecard_separates_assertions_from_guards(capsys):
    card = Scorecard()
    card.add(TaskResult("t1", "data_engineering", "analyze", "file", [
        CheckResult("a", True), CheckResult("b", False),
        CheckResult("g", False, "amount is text", category="guard")]))
    _print_scorecard(card)
    printed = capsys.readouterr().out
    lines = [" ".join(line.split()) for line in printed.splitlines()]
    assert "assertions 1/2 passed" in lines
    assert "guards 0/1 held" in lines
    assert "✗ [guard] g: amount is text" in lines
    assert "✗ b:" in lines


def test_a_guard_must_name_a_column_that_is_checked(tmp_path):
    assert _role_check_category(tmp_path, ("amount",)) == "guard"
    assert _role_check_category(tmp_path, ()) == "assertion"
    with pytest.raises(ValueError, match="guard_roles"):
        Expectations(expected_roles={"amount": "numeric"}, guard_roles=("amuont",))


def _role_check_category(out, guards):
    roles = {"amount": "numeric"}
    (out / "data_quality.json").write_text(json.dumps(
        {"columns": [{"name": name, "role": role} for name, role in roles.items()]}))
    (check,) = score_roles(out, Expectations(expected_roles=roles, guard_roles=guards))
    return check.category


# --- goal-analysis scorer ---------------------------------------------------
def _goal_answer(tmp_path, groups, status="answered"):
    (tmp_path / "goal_analysis.json").write_text(json.dumps({
        "kind": "group_comparison", "status": status,
        "metrics": {"group_column": "method", "value_column": "revenue",
                    "aggregation": "sum",
                    "groups": [{"group": g, "value": v} for g, v in groups.items()]}}))
    return tmp_path


def _by_name(checks):
    return {c.name: c for c in checks}


def test_goal_scorer_is_silent_when_nothing_is_expected(tmp_path):
    assert score_goal_analysis(tmp_path, Expectations()) == []


def test_goal_scorer_fails_when_the_artifact_is_missing(tmp_path):
    (check,) = score_goal_analysis(tmp_path, Expectations(expected_goal_status="limited"))
    assert not check.passed and "goal_analysis.json" in check.detail


def test_goal_scorer_grades_each_group_against_its_planted_value(tmp_path):
    exp = Expectations(expected_groups={"card": 1200.0, "None": 800.0})
    good = _by_name(score_goal_analysis(
        _goal_answer(tmp_path, {"card": 1200.0, "None": 800.0}), exp))
    assert set(good) == {"goal_group[card]", "goal_group[None]"}
    assert all(c.passed for c in good.values())
    wrong = _by_name(score_goal_analysis(
        _goal_answer(tmp_path, {"card": 1200.0, "None": 801.0}), exp))
    assert wrong["goal_group[card]"].passed and not wrong["goal_group[None]"].passed
    absent = _by_name(score_goal_analysis(_goal_answer(tmp_path, {"card": 1200.0}), exp))
    assert not absent["goal_group[None]"].passed
    assert "absent" in absent["goal_group[None]"].detail


def test_a_conserved_total_does_not_pass_the_goal_scorer(tmp_path):
    """The shape of the loss: three categories merged into "nan", total intact."""
    planted = {"card": 1200.0, "None": 800.0, "NA": 500.0, "N/A": 300.0}
    merged = {"card": 1200.0, "nan": 1600.0}
    assert sum(planted.values()) == sum(merged.values())
    checks = _by_name(score_goal_analysis(
        _goal_answer(tmp_path, merged),
        Expectations(expected_groups=planted, forbid_unexpected_groups=True)))
    assert checks["goal_group[card]"].passed
    assert not checks["goal_group[None]"].passed
    unexpected = checks["goal_groups_no_unexpected"]
    assert not unexpected.passed
    assert unexpected.detail == "2 groups, expected 4; unexpected group 'nan'=1600.0"


def test_forbid_unexpected_groups_is_what_refuses_an_extra_group(tmp_path):
    """Every expected group right, one group nobody planted: only the flag sees it."""
    out = _goal_answer(tmp_path, {"card": 1200.0, "nan": 1600.0})
    lenient = score_goal_analysis(out, Expectations(expected_groups={"card": 1200.0}))
    assert [c.name for c in lenient] == ["goal_group[card]"] and lenient[0].passed
    strict = _by_name(score_goal_analysis(out, Expectations(
        expected_groups={"card": 1200.0}, forbid_unexpected_groups=True)))
    assert not strict["goal_groups_no_unexpected"].passed
    exact = _by_name(score_goal_analysis(
        _goal_answer(tmp_path, {"card": 1200.0}),
        Expectations(expected_groups={"card": 1200.0}, forbid_unexpected_groups=True)))
    assert exact["goal_groups_no_unexpected"].passed


def test_goal_scorer_compares_the_status(tmp_path):
    out = _goal_answer(tmp_path, {"east": 20.0}, status="answered")
    (red,) = score_goal_analysis(out, Expectations(expected_goal_status="limited"))
    assert red.name == "goal_status" and not red.passed
    assert "status 'answered', expected 'limited'" in red.detail
    (green,) = score_goal_analysis(out, Expectations(expected_goal_status="answered"))
    assert green.passed


# --- profile fidelity -------------------------------------------------------
def _profile(tmp_path, column):
    (tmp_path / "data_quality.json").write_text(json.dumps(
        {"quality": {"score": 100}, "columns": [column]}))
    return tmp_path


def test_profile_scorer_checks_exact_label_counts_and_exact_gaps(tmp_path):
    exp = Expectations(expected_category_counts={"method": {"card": 2, "None": 1}},
                       expected_missing={"method": 0})
    kept = _by_name(score_quality(_profile(tmp_path, {
        "name": "method", "missing": 0, "n_unique": 2,
        "top_values": [{"value": "card", "count": 2}, {"value": "None", "count": 1}]}), exp))
    assert kept["category_counts[method]"].passed and kept["missing[method]"].passed
    lost = _by_name(score_quality(_profile(tmp_path, {
        "name": "method", "missing": 1, "n_unique": 1,
        "top_values": [{"value": "card", "count": 2}]}), exp))
    assert not lost["category_counts[method]"].passed
    assert lost["missing[method]"].detail == "method missing=1, expected 0"
    assert not lost["missing[method]"].passed
    extra = _by_name(score_quality(_profile(tmp_path, {
        "name": "method", "missing": 0, "n_unique": 3,
        "top_values": [{"value": "card", "count": 2}, {"value": "None", "count": 1},
                       {"value": "nan", "count": 1}]}), exp))
    assert not extra["category_counts[method]"].passed


# --- the planted dataset and its task ---------------------------------------
def test_null_labelled_payments_is_written_as_planted(tmp_path):
    path = ds.build_null_labelled_payments(tmp_path / "p.csv")
    raw = pd.read_csv(path, dtype=str, keep_default_na=False)   # the cells as written
    assert len(raw) == 55
    assert raw["payment_method"].value_counts().to_dict() == ds.NULL_LABELLED_PAYMENT_COUNTS
    assert {"None", "NA", "N/A"} <= set(raw["payment_method"])
    assert (raw["revenue"].astype(float) == 100.0).all()
    assert (raw["amount"] == "NA").sum() == 3
    assert pd.to_numeric(raw.loc[raw["amount"] != "NA", "amount"]).notna().all()
    with pytest.raises(ValueError, match="exact"):
        ds.build_null_labelled_payments(tmp_path / "big.csv", n=5000)


def test_null_word_task_runs_and_marks_exactly_one_guard(tmp_path):
    """The task grades the real mission; whether it is green is the suite's business."""
    (task,) = [t for t in data_engineering_tasks() if t.name == "de_null_words_stay_categories"]
    (result,) = BenchmarkRunner(tmp_path).run([task]).results
    assert result.error is None, result.error
    by_name = _by_name(result.checks)
    assert [c.name for c in result.checks if c.category == "guard"] == ["role[amount]"]
    assert by_name["role[amount]"].passed
    assert {"goal_groups_no_unexpected", "category_counts[payment_method]",
            "missing[payment_method]", "goal_group[None]", "goal_group[NA]",
            "goal_group[N/A]", "goal_group[card]"} <= set(by_name)
    # the total is conserved either way, which is why no check compares it
    answer = json.loads((tmp_path / task.name / "file" / "mission" / "goal_analysis.json").read_text())
    assert sum(g["value"] for g in answer["metrics"]["groups"]) == pytest.approx(5500.0)


# --- a sheet whose table does not start in A1 --------------------------------
def _schema_profile(tmp_path, names, n_rows=12):
    (tmp_path / "data_quality.json").write_text(json.dumps(
        {"dataset": {"n_rows": n_rows}, "columns": [{"name": n} for n in names]}))
    return tmp_path


def test_schema_scorer_is_silent_when_nothing_is_expected(tmp_path):
    assert score_schema(tmp_path, Expectations()) == []


def test_schema_scorer_checks_names_placeholders_and_rows(tmp_path):
    exp = Expectations(expected_columns=["region", "units"], forbid_unnamed_columns=True,
                       expected_n_rows=12)
    good = _by_name(score_schema(_schema_profile(tmp_path, ["region", "units"]), exp))
    assert set(good) == {"columns_exact", "no_unnamed_columns", "n_rows"}
    assert all(c.passed for c in good.values())
    indented = _by_name(score_schema(
        _schema_profile(tmp_path, ["Unnamed: 0", "region", "units"]), exp))
    assert not indented["columns_exact"].passed and not indented["no_unnamed_columns"].passed
    assert indented["no_unnamed_columns"].detail == "placeholder column names ['Unnamed: 0']"
    title_as_header = _by_name(score_schema(
        _schema_profile(tmp_path, ["Quarterly sales report", "Unnamed: 2"], n_rows=14), exp))
    assert not any(c.passed for c in title_as_header.values())
    reordered = _by_name(score_schema(_schema_profile(tmp_path, ["units", "region"]), exp))
    assert not reordered["columns_exact"].passed and reordered["no_unnamed_columns"].passed
    # a real column that merely contains the word is not a placeholder
    named = _by_name(score_schema(_schema_profile(tmp_path, ["Unnamed: north", "units"]),
                                  Expectations(forbid_unnamed_columns=True)))
    assert named["no_unnamed_columns"].passed


def test_guard_checks_mark_exactly_the_named_checks():
    checks = [CheckResult("columns_exact", True), CheckResult("economics_sane", True)]
    marked = mark_guards(checks, Expectations(guard_checks=("columns_exact",)))
    assert [(c.name, c.category) for c in marked] == [
        ("columns_exact", "guard"), ("economics_sane", "assertion")]


def test_a_guard_that_never_ran_is_an_error_not_a_pass():
    with pytest.raises(ValueError, match="matched no check"):
        mark_guards([CheckResult("economics_sane", True)],
                    Expectations(guard_checks=("columns_exact",)))


def test_titled_sheet_is_written_as_planted(tmp_path):
    import openpyxl          # not importorskip: a missing engine must fail, not skip

    sheet = openpyxl.load_workbook(ds.build_titled_sheet(tmp_path / "s.xlsx")).active
    assert sheet["B1"].value == ds.TITLED_SHEET_TITLE and sheet["A1"].value is None
    assert [c.value for c in sheet[2]] == [None, None, None, None]            # blank row
    assert [c.value for c in sheet[3]] == [None] + ds.TITLED_SHEET_HEADER       # indented header
    assert all(sheet.cell(row=r, column=1).value is None for r in range(1, 16))
    assert sheet.freeze_panes == "C4"
    assert sheet.max_row == 15                                                  # 12 data rows


def test_spreadsheet_header_task_holds_with_three_guards(tmp_path):
    (task,) = [t for t in data_engineering_tasks() if t.name == "de_spreadsheet_header_is_found"]
    (result,) = BenchmarkRunner(tmp_path).run([task]).results
    assert result.passed, result.as_dict()
    assert sorted(c.name for c in result.checks if c.category == "guard") == [
        "columns_exact", "n_rows", "no_unnamed_columns"]
    profile = json.loads((tmp_path / task.name / "file" / "mission" / "data_quality.json").read_text())
    assert [c["name"] for c in profile["columns"]] == ["region", "units", "revenue"]


def test_a_dataset_that_cannot_be_built_fails_and_is_never_a_skip(tmp_path, capsys):
    """Without openpyxl a spreadsheet task cannot build its file. That is a
    failed run and a non-zero exit, and the tasks after it still run."""
    def no_openpyxl(dest):
        raise ModuleNotFoundError("No module named 'openpyxl'")

    def driver_unavailable(dest):
        raise RuntimeError("driver unavailable")

    def task(name, builder):
        return BenchmarkTask(name=name, family="data_engineering", kind="analyze", goal="x",
                             dataset=DatasetSpec(name, "synthetic", builder=builder,
                                                 suffix=".xlsx"))
    after = BenchmarkTask(name="after", family="data_engineering", kind="analyze", goal="x",
                          dataset=DatasetSpec("churn", "synthetic", builder=ds.build_churn))
    card = BenchmarkRunner(tmp_path).run(
        [task("no_engine", no_openpyxl), task("worded", driver_unavailable), after])
    first, second, third = card.results
    assert not first.passed and "openpyxl" in first.error
    assert not second.passed
    assert third.passed                                   # the suite carried on
    # the exit rule in __main__ skips an error containing this one word
    assert "unavailable" not in first.error and "unavailable" not in second.error
    _print_scorecard(card)
    printed = capsys.readouterr().out
    assert printed.count("[FAIL]") == 2 and "[SKIP]" not in printed
    assert "dataset could not be built: ModuleNotFoundError" in printed


# --- a builder that raises: one failed run, a non-zero exit, never a pass -----
def _raising(exc):
    def builder(dest):
        raise exc
    return builder


def _unbuildable(name, exc):
    return BenchmarkTask(name=name, family="data_engineering", kind="analyze", goal="x",
                         dataset=DatasetSpec(name, "synthetic", builder=_raising(exc)))


def test_a_builder_that_raises_is_a_failed_run_and_the_command_exits_1(tmp_path, monkeypatch):
    """Every way a broken run could be mistaken for a good one, at once."""
    from avaloka.benchmark import __main__ as entry

    good = BenchmarkTask(name="good", family="data_engineering", kind="analyze", goal="x",
                         dataset=DatasetSpec("churn", "synthetic", builder=ds.build_churn))
    tasks = [_unbuildable("broken", RuntimeError("boom")), good]
    card = BenchmarkRunner(tmp_path / "direct").run(tasks)
    broken, fine = card.results
    assert broken.error and "boom" in broken.error            # its error is set
    assert broken.passed is False                              # its passed is False
    assert broken in card.failures()                           # it is in failures()
    assert fine.passed and fine not in card.failures()         # the next task still ran
    counts = card.as_dict()
    assert counts["n_runs"] == 2 and counts["n_passed"] == 1   # not counted as passed

    # and through the command: exit code 1, scorecard written with the same counts
    monkeypatch.setattr(entry, "all_tasks",
                        lambda include_kaggle=False, include_edgar=False: tasks)
    out_json = tmp_path / "card.json"
    # --family pins this to the two tasks above. `main()` appends the
    # conversational family to every unfiltered run, so the scorecard totals
    # asserted below would otherwise count six scorers this test never set up.
    monkeypatch.setattr("sys.argv", ["avaloka.benchmark", "run", "--quiet",
                                     "--family", "data_engineering",
                                     "--out", str(tmp_path / "cmd"), "--json", str(out_json)])
    with pytest.raises(SystemExit) as stop:
        entry.main()
    assert stop.value.code == 1
    written = json.loads(out_json.read_text())
    assert written["n_runs"] == 2 and written["n_passed"] == 1
    assert written["results"][0]["passed"] is False and written["results"][0]["error"]


@pytest.mark.parametrize("exc", [
    RuntimeError("boom"), ModuleNotFoundError("No module named 'openpyxl'"),
    ValueError(""), OSError("disk full"), RuntimeError("driver unavailable"), KeyError("col"),
], ids=lambda e: type(e).__name__ + (":" + str(e) if str(e) else ":empty"))
def test_a_builder_failure_never_yields_a_result_with_no_error_and_no_checks(tmp_path, exc):
    """The hole: TaskResult.passed is `error is None and all(checks)`, so a
    result with no checks and no error reads as PASSED. A builder failure must
    never produce one, whatever was raised and however empty its message."""
    (result,) = BenchmarkRunner(tmp_path).run([_unbuildable("broken", exc)]).results
    assert not (result.error is None and not result.checks)
    assert result.error and result.passed is False
    assert "unavailable" not in result.error                   # never read as a skip
    # the shape that must not come back, shown to be the dangerous one
    assert TaskResult("t", "data_engineering", "analyze", "—", checks=[], error=None).passed


def test_a_missing_spreadsheet_library_names_the_extra_to_install(tmp_path, capsys):
    import tomllib
    from pathlib import Path

    from avaloka.benchmark.runner import BENCHMARK_EXTRA_HINT

    task = _unbuildable("sheet", ModuleNotFoundError("No module named 'openpyxl'"))
    (result,) = (card := BenchmarkRunner(tmp_path).run([task])).results
    assert BENCHMARK_EXTRA_HINT == 'pip install "avaloka[benchmark]"'
    assert BENCHMARK_EXTRA_HINT in result.error
    _print_scorecard(card)
    assert BENCHMARK_EXTRA_HINT in capsys.readouterr().out
    # the extra the message names exists, holds the library, and is not a runtime dependency
    project = tomllib.loads(
        (Path(__file__).resolve().parents[2] / "pyproject.toml").read_text())["project"]
    assert project["name"] == "avaloka"
    assert any(d.startswith("openpyxl") for d in project["optional-dependencies"]["benchmark"])
    assert not any(d.lower().startswith("openpyxl") for d in project["dependencies"])
    assert len(project["dependencies"]) == 8
